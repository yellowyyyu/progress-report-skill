# -*- coding: utf-8 -*-
"""进展汇报 Word 生成模板（依赖 python-docx）。

用法：
  1. 改下面的 OUT 为输出路径
  2. 在"内容区"按说明填内容：
     - heading(doc, "10月2日") / heading(doc, "Model-v12/v13")   → 日期行、主题行（加粗）
     - para(doc, [("问题：", {"bold": True}), "……"])              → 标签加粗的正文段
     - table(doc, ["对比项", "旧", "新"], [["…", "…", "…"]])      → 最简网格对照表
     - figure(doc, "x.png", "图1　说明")                          → 居中插图＋图注
     - blank(doc)                                                 → 块与块之间的空行（必须留）
  3. 运行：python make_report_docx.py

排版（已预设，勿随意改）：宋体＋Times New Roman、五号(10.5pt)、单倍行距、
首行缩进两字符、两端对齐、黑色、仅标题行加粗。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path("进展汇报.docx")  # ← 改成你的输出路径

SIZE = 10.5           # 五号
INDENT = Pt(21)       # 首行缩进两字符（2 × 10.5pt）
CN_FONT = "宋体"
EN_FONT = "Times New Roman"


def set_font(run, bold=False):
    run.font.name = EN_FONT
    run.font.size = Pt(SIZE)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CN_FONT)


def fmt(p, indent=True, align=None):
    pf = p.paragraph_format
    pf.first_line_indent = INDENT if indent else Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY


def para(doc, segments, indent=True, align=None):
    """segments: 字符串或 (文本, {"bold": True/False}) 组成的列表"""
    p = doc.add_paragraph()
    fmt(p, indent, align)
    for seg in segments:
        if isinstance(seg, str):
            seg = (seg, {})
        text, opts = seg
        r = p.add_run(text)
        set_font(r, bold=opts.get("bold", False))
    return p


def heading(doc, text):
    """日期行 / 主题行（加粗、不缩进）"""
    para(doc, [(text, {"bold": True})], indent=False)


def table(doc, header, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_font(r, bold=True)
        fmt(cell.paragraphs[0], indent=False)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(v)
            set_font(r)
            fmt(cell.paragraphs[0], indent=False)
    return t


def figure(doc, path, caption, width_cm=14.5):
    p = doc.add_paragraph()
    fmt(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    fmt(cap, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = cap.add_run(caption)
    set_font(r)
    return cap


def blank(doc):
    """块与块之间的空行"""
    p = doc.add_paragraph()
    fmt(p, indent=False)
    return p


# ================= 内容区（从这里开始改） =================

doc = Document()
para(doc, [("××进展汇报（X月X日—X月X日）", {"bold": True})],
     indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc, [("总览：", {"bold": True}),
           "一句话项目定位。现在能做到××（关键数字）。本期做了××。当前状态××。"])
blank(doc)

heading(doc, "X月X日")
heading(doc, "阶段主题一")
para(doc, [("问题：", {"bold": True}), "大白话讲清发现了什么问题，附关键数字。"])
para(doc, [("做法：", {"bold": True}), "1.专业术语标签：白话解释。2.……"])
para(doc, [("结果：", {"bold": True}), "关键数字，如实写短板。"])
blank(doc)

heading(doc, "X月X日")
heading(doc, "阶段主题二")
para(doc, [("问题：", {"bold": True}), "……"])
para(doc, [("做法：", {"bold": True}), "……"])
table(doc, ["对比项", "旧", "新"], [
    ["指标一", "旧值", "新值"],
    ["指标二", "旧值", "新值"],
])
blank(doc)
para(doc, [("结果：", {"bold": True}), "……"])
blank(doc)

para(doc, [("后续，××。", {"bold": True})])

doc.save(OUT)
print("saved", OUT)
